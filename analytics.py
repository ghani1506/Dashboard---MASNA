import io, math, collections
import openpyxl

GRADES=['A*','A2','B3','B4','C5','C6','D7','E8','U']
DEFAULT_BOUNDS=[89.5,79.5,74.5,69.5,64.5,59.5,49.5,39.5,0]
def valid(x): return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x) and 0<=x<=100
def grade(x,bounds=DEFAULT_BOUNDS):
 if not valid(x): return None
 return next((g for g,b in zip(GRADES,bounds) if x>=b),'U')
def parse_workbook(data,name):
 w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
 try:
  key=next((s for s in w.sheetnames if s.lower()=='master'),None)
  if key is None: raise ValueError('Master sheet not found')
  a=list(w[key].values)
  def v(r,c): return a[r-1][c-1] if r<=len(a) and c<=len(a[r-1]) else None
  heads=[(c,v(10,c)) for c in range(8,50) if isinstance(v(10,c),str) and not any(t in v(10,c).upper() for t in ['MARK','POSITION']) and v(10,c)!='BUSINESS, ART & TECHNOLOGY']
  records=[]
  for r in range(14,44):
   if not v(r,2): continue
   marks=[]
   for c,s in heads:
    raw=v(r,c+2)
    marks.append(dict(subject=s,score=raw if valid(raw) else None,task=v(r,c),continuous=v(r,c+1),status='Assessed' if valid(raw) else 'Missing / not entered' if raw in (None,'') else str(raw),sourceCell=openpyxl.utils.get_column_letter(c+2)+str(r)))
   records.append(dict(id=name+':'+str(r),name=str(v(r,2)),row=r,marks=marks))
  if not records: raise ValueError('No student records in B14:B43')
  errors=collections.Counter(x for row in a for x in row if isinstance(x,str) and x in ['#REF!','#DIV/0!','#N/A','#VALUE!'])
  return dict(file=name,year=str(v(3,1)),exam=str(v(4,1)),**{'class':str(v(7,1))},records=records,errors=dict(errors))
 finally: w.close()
def flatten(books,component,bounds,excluded):
 rows=[]
 for b in books:
  for r in b['records']:
   for m in r['marks']:
    x=m.get(component); x=x if valid(x) else None
    rows.append({'ID':r['id'],'Student':r['name'],'Class':b['class'],'Year':b['year'],'Examination':b['exam'],'Subject':m['subject'],'Mark':x,'Grade':grade(x,bounds),'Eligible':m['subject'] not in excluded,'Status':'Assessed' if x is not None else m.get('status','Missing'),'Source':b['file'],'Source cell':m.get('sourceCell','')})
 return rows
